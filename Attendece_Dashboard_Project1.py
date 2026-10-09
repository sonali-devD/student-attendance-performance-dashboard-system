import os
import tkinter as tk
from tkinter import messagebox, ttk
import openpyxl
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg

EXCEL_FILE = r"C:\Users\Sonal\Downloads\student_records.xlsx"

if not os.path.exists(EXCEL_FILE):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(["Roll No", "Name", "Class", "Java", "Python", "AI", "PHP",
               "HTML", "Attendance (%)", "Total", "Percentage", "Result"])
    wb.save(EXCEL_FILE)


def open_summary_window(s_name, r_no, attendance_val, total, percentage, result):
    sum_win = tk.Toplevel(window)
    sum_win.title("Student Record Summary")
    sum_win.geometry("600x550")
    sum_win.config(bg="#fef3c7")

    tk.Label(sum_win, text="Record Saved Successfully!", font=("Arial", 20, "bold"), bg="#fef3c7", fg="#15803d").pack(pady=15)
 
    summary_text = (
        f"Student Name : {s_name}\n"
        f"Roll Number  : {r_no}\n"
        f"Attendance   : {attendance_val}%\n"
        f"Total Marks  : {total} / 500\n"
        f"Percentage   : {percentage:.2f}%\n"
        f"Final Result : {result}"
    )
    
    lbl = tk.Label(sum_win, text=summary_text, font=("Arial", 14, "bold"), bg="white", fg="#1e293b", justify="left", bd=2, relief="groove", padx=15, pady=15)
    lbl.pack(pady=10, padx=20, fill="x")

    tk.Button(sum_win, text="Close Window", bg="#dc2626", fg="white", font=("Arial", 10, "bold"), command=sum_win.destroy).pack(pady=10)


def open_save_form():
    form_win = tk.Toplevel(window)
    form_win.title("Add New Student Record")
    form_win.geometry("900x650")
    form_win.config(bg="#fbcfe8")

    tk.Label(form_win, text="Enter Student Details", font=("Georgia", 20, "bold"), bg="lavender", fg="#121212", height=2).pack(pady=10)
    fields_frame = tk.Frame(form_win, bg="#fbcfe8")
    fields_frame.pack(pady=5)

    labels = ["Name:", "Roll No:", "Class:", "Java Marks:", "Python Marks:", "AI Marks:", "PHP Marks:", "HTML Marks:", "Attendance (%):"]
    entries = []

    for idx, text in enumerate(labels):
        tk.Label(fields_frame, text=text, bg="#fbcfe8", font=("Arial", 15, "bold")).grid(row=idx, column=0, sticky="e", pady=3, padx=5)
        e = tk.Entry(fields_frame, width=22, font=("Arial", 14, "bold"))
        e.grid(row=idx, column=1, pady=3, padx=5)
        entries.append(e)

    name_e, roll_e, class_e, sub1_e, sub2_e, sub3_e, sub4_e, sub5_e, att_e = entries

    def submit_data():
        s_name = name_e.get().strip()
        s_roll = roll_e.get().strip()
        s_class = class_e.get().strip()
        p1 = sub1_e.get().strip()
        p2 = sub2_e.get().strip()
        p3 = sub3_e.get().strip()
        p4 = sub4_e.get().strip()
        p5 = sub5_e.get().strip()
        att = att_e.get().strip()

        if not (s_name and s_roll and s_class and p1 and p2 and p3 and p4 and p5 and att):
            messagebox.showerror("Error", "Please fill all fields, including attendance!")
            return

        try:
            r_no = int(s_roll)
            m1, m2, m3, m4, m5 = float(p1), float(p2), float(p3), float(p4), float(p5)
            attendance_val = float(att)
            if not (0 <= attendance_val <= 100):
                raise ValueError
        except ValueError:
            messagebox.showerror("Error", "Roll No, Marks, and Attendance must be valid numbers (Attendance: 0-100)!")
            return

        total = m1 + m2 + m3 + m4 + m5
        percentage = total / 5.0

        if m1 >= 35 and m2 >= 35 and m3 >= 35 and m4 >= 35 and m5 >= 35:
            result = "Pass"
        else:
            result = "Fail"

        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active

        for row in ws.iter_rows(min_row=2, values_only=True):
            if str(row[0]) == str(s_roll):
                messagebox.showerror("Error", "Roll Number already exists!")
                wb.close()
                return

        ws.append([r_no, s_name, s_class, m1, m2, m3, m4, m5, attendance_val, total, round(percentage, 2), result])
        wb.save(EXCEL_FILE)
        wb.close()

        form_win.destroy()

        if messagebox.showinfo("Success", "Record Saved Successfully!") == "ok":
            open_summary_window(s_name, r_no, attendance_val, total, percentage, result)

    tk.Button(form_win, text="Save Record", width=30, bg="#16a34a", fg="white", font=("Arial", 15, "bold"), command=submit_data).pack(pady=15)


def open_all_records():
    rec_win = tk.Toplevel(window)
    rec_win.title("All Student Records")
    rec_win.geometry("950x650")
    rec_win.config(bg="#cffafe")

    tk.Label(rec_win, text="Registered Student Records Database", font=("Arial", 14, "bold"), bg="#cffafe", fg="#0369a1").pack(pady=10)

    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
        data = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]
        wb.close()
    except Exception:
        data = []

    if not data:
        tk.Label(rec_win, text="No records found in Excel file.", bg="#cffafe", font=("Arial", 11)).pack(pady=20)
        return

    # Sort data Roll No wise (ascending order)
    data = sorted(data, key=lambda x: int(x[0]) if str(x[0]).isdigit() else str(x[0]))

    columns = ("Roll No", "Name", "Class", "Java", "Python", "AI", "PHP", "HTML", "Attendance", "Total", "%", "Result")
    tree = ttk.Treeview(rec_win, columns=columns, show="headings", height=14)

    for col in columns:
        tree.heading(col, text=col)
        tree.column(col, width=75, anchor="center")

    for row in data:
        tree.insert("", tk.END, values=(row[0], row[1], row[2], row[3], row[4], row[5], row[6], row[7], f"{row[8]}%", row[9], f"{row[10]}%", row[11]))

    tree.pack(fill="both", expand=True, padx=15, pady=10)

    # --- UPDATE RECORD FUNCTION ---
    def update_selected_record():
        selected = tree.selection()
        if not selected:
            messagebox.showwarning("Warning", "Please select a student record to update!")
            return
        
        item = tree.item(selected)
        roll_no = item['values'][0]
        
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
        
        row_data = None
        row_idx = None
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if str(row[0]) == str(roll_no):
                row_data = row
                row_idx = idx
                break
        wb.close()
        
        if not row_data:
            messagebox.showerror("Error", "Record not found in Excel!")
            return
        
        up_win = tk.Toplevel(rec_win)
        up_win.title(f"Update Record - Roll No: {roll_no}")
        up_win.geometry("900x650")
        up_win.config(bg="#fbcfe8")
        
        tk.Label(up_win, text=f"Update Details for Roll No: {roll_no}", font=("Georgia", 18, "bold"), bg="lavender", fg="#121212", height=2).pack(pady=10)
        
        fields_frame = tk.Frame(up_win, bg="#fbcfe8")
        fields_frame.pack(pady=5)
        
        labels = ["Name:", "Roll No:", "Class:", "Java Marks:", "Python Marks:", "AI Marks:", "PHP Marks:", "HTML Marks:", "Attendance (%):"]
        entries = []
        default_vals = [row_data[1], row_data[0], row_data[2], row_data[3], row_data[4], row_data[5], row_data[6], row_data[7], row_data[8]]
        
        for idx, text in enumerate(labels):
            tk.Label(fields_frame, text=text, bg="#fbcfe8", font=("Arial", 15, "bold")).grid(row=idx, column=0, sticky="e", pady=3, padx=5)
            e = tk.Entry(fields_frame, width=22, font=("Arial", 14, "bold"))
            e.insert(0, str(default_vals[idx]))
            e.grid(row=idx, column=1, pady=3, padx=5)
            entries.append(e)
        
        name_e, roll_e, class_e, sub1_e, sub2_e, sub3_e, sub4_e, sub5_e, att_e = entries
        
        def save_updated_data():
            s_name = name_e.get().strip()
            s_roll = roll_e.get().strip()
            s_class = class_e.get().strip()
            p1 = sub1_e.get().strip()
            p2 = sub2_e.get().strip()
            p3 = sub3_e.get().strip()
            p4 = sub4_e.get().strip()
            p5 = sub5_e.get().strip()
            att = att_e.get().strip()
            
            if not (s_name and s_roll and s_class and p1 and p2 and p3 and p4 and p5 and att):
                messagebox.showerror("Error", "Please fill all fields!")
                return
            
            try:
                r_no = int(s_roll)
                m1, m2, m3, m4, m5 = float(p1), float(p2), float(p3), float(p4), float(p5)
                attendance_val = float(att)
                if not (0 <= attendance_val <= 100):
                    raise ValueError
            except ValueError:
                messagebox.showerror("Error", "Roll No, Marks, and Attendance must be valid numbers (Attendance: 0-100)!")
                return
            
            total = m1 + m2 + m3 + m4 + m5
            percentage = total / 5.0
            result = "Pass" if all(x >= 35 for x in [m1, m2, m3, m4, m5]) else "Fail"
            
            wb = openpyxl.load_workbook(EXCEL_FILE)
            ws = wb.active
            
            if str(r_no) != str(roll_no):
                for r in ws.iter_rows(min_row=2, values_only=True):
                    if str(r[0]) == str(r_no):
                        messagebox.showerror("Error", "New Roll Number already exists!")
                        wb.close()
                        return
            
            ws.cell(row=row_idx, column=1, value=r_no)
            ws.cell(row=row_idx, column=2, value=s_name)
            ws.cell(row=row_idx, column=3, value=s_class)
            ws.cell(row=row_idx, column=4, value=m1)
            ws.cell(row=row_idx, column=5, value=m2)
            ws.cell(row=row_idx, column=6, value=m3)
            ws.cell(row=row_idx, column=7, value=m4)
            ws.cell(row=row_idx, column=8, value=m5)
            ws.cell(row=row_idx, column=9, value=attendance_val)
            ws.cell(row=row_idx, column=10, value=total)
            ws.cell(row=row_idx, column=11, value=round(percentage, 2))
            ws.cell(row=row_idx, column=12, value=result)
            
            wb.save(EXCEL_FILE)
            wb.close()
            
            up_win.destroy()
            rec_win.destroy()
            messagebox.showinfo("Success", "Record Updated Successfully!")
            open_all_records()
        
        tk.Button(up_win, text="Save Changes", width=30, bg="#16a34a", fg="white", font=("Arial", 15, "bold"), command=save_updated_data).pack(pady=15)

    # --- DELETE RECORD FUNCTION ---
    def delete_selected_record():
        selected = tree.selection()
        if not selected:
            messagebox.showwarning("Warning", "Please select a student record to delete!")
            return
        
        item = tree.item(selected)
        roll_no = item['values'][0]
        
        if messagebox.askyesno("Confirm Delete", f"Are you sure you want to delete Roll No {roll_no}?"):
            wb = openpyxl.load_workbook(EXCEL_FILE)
            ws = wb.active
            
            row_to_delete = None
            for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if str(row[0]) == str(roll_no):
                    row_to_delete = idx
                    break
            
            if row_to_delete:
                ws.delete_rows(row_to_delete)
                wb.save(EXCEL_FILE)
                wb.close()
                messagebox.showinfo("Success", "Record deleted successfully!")
                rec_win.destroy()
                open_all_records()
            else:
                wb.close()
                messagebox.showerror("Error", "Record not found in Excel file.")

    btn_frame = tk.Frame(rec_win, bg="#cffafe")
    btn_frame.pack(pady=10)

    tk.Button(btn_frame, text="Update Selected Record", bg="#d97706", fg="white", font=("Arial", 11, "bold"), width=22, command=update_selected_record).pack(side=tk.LEFT, padx=10)
    tk.Button(btn_frame, text="Delete Selected Record", bg="#dc2626", fg="white", font=("Arial", 11, "bold"), width=22, command=delete_selected_record).pack(side=tk.LEFT, padx=10)

# --- VIEW ANALYTICS DASHBOARD WINDOW (MATCHING REFERENCE IMAGE) ---
def open_analytics():
    ana_win = tk.Toplevel(window)
    ana_win.title("student marks dashboard")
    ana_win.geometry("950x700")
    ana_win.config(bg="#ffffff")

    try:
        df = pd.read_excel(EXCEL_FILE)
        df.columns = df.columns.str.strip()
    except Exception as e:
        tk.Label(ana_win, text=f"Error reading data file: {e}", bg="#ffffff", fg="red").pack(pady=20)
        return

    if df.empty or len(df) == 0:
        tk.Label(ana_win, text="No student records available for analysis yet.", bg="#ffffff", font=("Arial", 11)).pack(pady=20)
        return

    fig, ax = plt.subplots(2, 2, figsize=(9, 6))
    
    subs = ["java", "python", "AI", "PHP", "HTML"]
    # Retrieve average marks for the 5 subjects
    avgs = [df["Java"].mean(), df["Python"].mean(), df["AI"].mean(), df["PHP"].mean(), df["HTML"].mean()]
    
    # 1. Top-Left: Bar chart (Subject wise Marks)
    ax[0, 0].bar(subs, avging := avgs)
    ax[0, 0].set_title("Subject wise Marks", fontsize=12)
    ax[0, 0].set_ylabel("marks", fontsize=10)
    ax[0, 0].set_xlabel("Subject", fontsize=10)
    ax[0, 0].set_ylim(0, 105)
    ax[0, 0].tick_params(axis="x", rotation=25)

    # 2. Top-Right: Line chart (marks trend)
    ax[0, 1].plot(subs, avgs, marker='o', linestyle='-')
    ax[0, 1].set_title("marks trend", fontsize=12)
    ax[0, 1].set_ylabel("marks", fontsize=10)
    ax[0, 1].set_xlabel("Subject", fontsize=10)
    ax[0, 1].set_ylim(78, 100)
    ax[0, 1].tick_params(axis="x", rotation=25)

    # 3. Bottom-Left: Pie chart (Subject share)
    ax[1, 0].pie(avgs, labels=subs, autopct='%1.1f%%', startangle=0)

    # 4. Bottom-Right: Summary Metrics Text Card
    overall_avg = sum(avgs) / len(avgs) if avgs else 0
    highest_mark = max(avgs) if avgs else 0
    lowest_mark = min(avgs) if avgs else 0
    
    summary = f"""STUDENT PERFORMANCE

Average Marks :{overall_avg:.2f}
Highest Marks :{highest_mark:.0f}
Lowest Marks :{lowest_mark:.0f}
Total Subjects :5"""

    ax[1, 1].axis("off")
    ax[1, 1].text(0.1, 0.3, summary, fontsize=13, verticalalignment="center", fontweight="normal")

    plt.suptitle("student marks dashboard", fontsize=16)
    plt.tight_layout()

    canvas = FigureCanvasTkAgg(fig, master=ana_win)
    canvas.draw()
    canvas.get_tk_widget().pack(fill="both", expand=True, padx=10, pady=10)


# --- MAIN WELCOME/HOME WINDOW ---
window = tk.Tk()
window.title("Student Attendance & Performance System - Welcome")
window.geometry("900x600")
window.config(bg="#e0e7ff")

header = tk.Label(window, text="Welcome to Student Attendence And Performance Management System", font=("Georgia", 15, "bold"), bg="#4f46e5", fg="white", height=2)
header.pack(fill="x", pady=10)

sub_lbl = tk.Label(window, text="Please choose an option below:", font=("Arial", 12, "bold"), bg="#e0e7ff", fg="#312e81")
sub_lbl.pack(pady=10)

menu_frame = tk.Frame(window, bg="#e0e7ff")
menu_frame.pack(pady=15)

tk.Button(menu_frame, text="Add / Save New Record", width=25, bg="#22c55e", fg="white", font=("Arial", 11, "bold"), pady=8, command=open_save_form).pack(pady=10)
tk.Button(menu_frame, text="View All Records", width=25, bg="#0ea5e9", fg="white", font=("Arial", 11, "bold"), pady=8, command=open_all_records).pack(pady=10)
tk.Button(menu_frame, text="View Analytics Dashboard", width=25, bg="#eab308", fg="white", font=("Arial", 11, "bold"), pady=8, command=open_analytics).pack(pady=10)

window.mainloop()
